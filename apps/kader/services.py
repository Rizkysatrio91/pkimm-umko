"""Service untuk app kader: bulk-import data kader dari Excel."""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO
from typing import Iterable

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.kader.models import BiodataKader
from core.constants import JabatanKader, JenisKelamin


# Kolom template Excel. Urutan ini WAJIB cocok dengan baris header di workbook.
EXCEL_COLUMNS = [
    ('npm', 'NPM', True),
    ('nama_lengkap', 'Nama Lengkap', True),
    ('program_studi', 'Program Studi', True),
    ('fakultas', 'Fakultas', True),
    ('angkatan', 'Angkatan (4 digit tahun)', True),
    ('jenis_kelamin', 'Jenis Kelamin (L/P)', False),
    ('tempat_lahir', 'Tempat Lahir', False),
    ('tgl_lahir', 'Tanggal Lahir (YYYY-MM-DD)', False),
    ('alamat_domisili', 'Alamat Domisili', False),
    ('no_hp', 'No HP', False),
    ('nama_ayah', 'Nama Ayah', False),
    ('nama_ibu', 'Nama Ibu', False),
    ('no_hp_ortu', 'No HP Orang Tua', False),
    ('jabatan_struktural', 'Jabatan Struktural (kode, lihat petunjuk)', False),
    ('periode_jabatan', 'Periode Jabatan (contoh: 2024-2025)', False),
]

VALID_JABATAN = {choice[0] for choice in JabatanKader.choices}
VALID_JK = {choice[0] for choice in JenisKelamin.choices}


@dataclass
class ImportResult:
    """Hasil bulk-import kader dari Excel."""
    created: int = 0
    updated: int = 0
    errors: list[tuple[int, str]] = field(default_factory=list)  # (row_number, pesan)

    @property
    def total_error(self) -> int:
        return len(self.errors)

    @property
    def has_error(self) -> bool:
        return bool(self.errors)


class KaderExcelService:
    """Generate template dan parse upload Excel data kader."""

    # ─── Generate template ──────────────────────────────────────────────────
    @staticmethod
    def generate_template() -> bytes:
        """Buat workbook Excel kosong dengan header, contoh 1 baris, dan styling."""
        wb = Workbook()
        ws = wb.active
        ws.title = 'Data Kader'

        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='5B1525', end_color='5B1525', fill_type='solid')
        center = Alignment(horizontal='center', vertical='center', wrap_text=True)

        # Header row
        for idx, (_, label, required) in enumerate(EXCEL_COLUMNS, start=1):
            cell = ws.cell(row=1, column=idx, value=f"{label}{' *' if required else ''}")
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            ws.column_dimensions[get_column_letter(idx)].width = max(18, len(label) + 4)
        ws.row_dimensions[1].height = 36

        # Contoh baris
        contoh = {
            'npm': '2259201005',
            'nama_lengkap': 'Contoh Nama Lengkap',
            'program_studi': 'Teknik Informatika',
            'fakultas': 'Fakultas Teknik dan Ilmu Komputer',
            'angkatan': '2022',
            'jenis_kelamin': 'L',
            'tempat_lahir': 'Kotabumi',
            'tgl_lahir': '2004-05-15',
            'alamat_domisili': 'Jl. Contoh No. 1',
            'no_hp': '081234567890',
            'nama_ayah': 'Nama Ayah',
            'nama_ibu': 'Nama Ibu',
            'no_hp_ortu': '081234567891',
            'jabatan_struktural': 'kader',
            'periode_jabatan': '2024-2025',
        }
        for idx, (key, _, _) in enumerate(EXCEL_COLUMNS, start=1):
            ws.cell(row=2, column=idx, value=contoh[key])

        # Keterangan singkat
        ws.cell(row=4, column=1, value='Petunjuk:').font = Font(bold=True, color='5B1525')
        ws.cell(row=5, column=1, value='- Kolom dengan tanda * wajib diisi.')
        ws.cell(
            row=6, column=1,
            value='- NPM adalah unik; jika sudah ada, baris akan di-update (bukan duplikat).',
        )
        ws.cell(
            row=7, column=1,
            value='- Jenis Kelamin gunakan huruf L atau P. Tanggal lahir format YYYY-MM-DD.',
        )
        ws.cell(
            row=8, column=1,
            value=f'- Jabatan yang valid: {", ".join(sorted(VALID_JABATAN))}.',
        )
        ws.cell(row=9, column=1, value='- Hapus baris contoh sebelum upload.')

        # Freeze header
        ws.freeze_panes = 'A2'

        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()

    # ─── Parse & import ─────────────────────────────────────────────────────
    @classmethod
    def import_from_file(cls, upload_file) -> ImportResult:
        """Baca workbook upload, validasi, dan simpan data.

        Satu baris gagal tidak menggagalkan baris lainnya — setiap baris pakai
        savepoint sendiri sehingga error di-skip.
        """
        try:
            wb = load_workbook(upload_file, data_only=True)
        except Exception as e:
            return ImportResult(errors=[(0, f'File Excel tidak dapat dibaca: {e}')])

        ws = wb.active
        result = ImportResult()
        rows = ws.iter_rows(min_row=2, values_only=True)

        for row_idx, row in enumerate(rows, start=2):
            # Skip baris kosong penuh
            if row is None or all(cell in (None, '') for cell in row):
                continue
            # Skip bagian "Petunjuk" di template bawaan (kolom B-... kosong, hanya A terisi teks non-data)
            if row[0] and isinstance(row[0], str) and row[0].strip().startswith('-'):
                continue
            if row[0] == 'Petunjuk:':
                continue

            try:
                with transaction.atomic():
                    created = cls._process_row(row)
                if created:
                    result.created += 1
                else:
                    result.updated += 1
            except ValueError as e:
                result.errors.append((row_idx, str(e)))
            except Exception as e:
                result.errors.append((row_idx, f'Error tidak terduga: {e}'))

        return result

    @classmethod
    def _process_row(cls, row: Iterable) -> bool:
        """Proses satu baris. Return True kalau insert baru, False kalau update.

        Raises ValueError jika data tidak valid.
        """
        cells = list(row) + [None] * (len(EXCEL_COLUMNS) - len(row))
        data = {key: cells[idx] for idx, (key, _, _) in enumerate(EXCEL_COLUMNS)}

        # Required fields
        for key, label, required in EXCEL_COLUMNS:
            if required and (data[key] is None or str(data[key]).strip() == ''):
                raise ValueError(f'Kolom wajib "{label}" kosong.')

        # Normalisasi string
        for key in data:
            if isinstance(data[key], str):
                data[key] = data[key].strip()

        npm = str(data['npm']).strip()
        if not npm.isdigit() or len(npm) < 7:
            raise ValueError(f'NPM "{npm}" tidak valid (harus angka minimal 7 digit).')

        # Angkatan: 4 digit
        angkatan = str(data['angkatan']).strip()
        if angkatan.endswith('.0'):
            angkatan = angkatan[:-2]
        if not (angkatan.isdigit() and len(angkatan) == 4):
            raise ValueError(f'Angkatan "{data["angkatan"]}" harus 4 digit tahun.')

        # Jenis kelamin
        jk = (data.get('jenis_kelamin') or '').strip().upper()[:1]
        if jk and jk not in VALID_JK:
            raise ValueError(f'Jenis kelamin "{data["jenis_kelamin"]}" harus L atau P.')

        # Tanggal lahir
        tgl_lahir = cls._parse_date(data.get('tgl_lahir'))

        # Jabatan
        jabatan = (data.get('jabatan_struktural') or 'calon').strip().lower()
        if jabatan not in VALID_JABATAN:
            raise ValueError(
                f'Jabatan "{jabatan}" tidak valid. Pilihan: {sorted(VALID_JABATAN)}.'
            )

        fields_map = {
            'nama_lengkap': str(data['nama_lengkap']).strip()[:100],
            'program_studi': str(data['program_studi']).strip()[:50],
            'fakultas': str(data['fakultas']).strip()[:50],
            'angkatan': angkatan,
            'tempat_lahir': (str(data.get('tempat_lahir') or '').strip())[:50],
            'tgl_lahir': tgl_lahir,
            'jenis_kelamin': jk,
            'alamat_domisili': str(data.get('alamat_domisili') or '').strip(),
            'no_hp': (str(data.get('no_hp') or '').strip())[:15],
            'nama_ayah': (str(data.get('nama_ayah') or '').strip())[:100],
            'nama_ibu': (str(data.get('nama_ibu') or '').strip())[:100],
            'no_hp_ortu': (str(data.get('no_hp_ortu') or '').strip())[:15],
            'jabatan_struktural': jabatan,
            'periode_jabatan': (str(data.get('periode_jabatan') or '').strip())[:20],
        }

        biodata = BiodataKader.objects.filter(npm=npm).first()
        if biodata:
            for k, v in fields_map.items():
                setattr(biodata, k, v)
            biodata.save()
            biodata.check_completeness()
            return False

        BiodataKader.objects.create(npm=npm, **fields_map).check_completeness()
        return True

    @staticmethod
    def _parse_date(value) -> date | None:
        if value is None or value == '':
            return None
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        s = str(value).strip()[:10]
        for fmt in ('%Y-%m-%d', '%d-%m-%Y', '%d/%m/%Y', '%Y/%m/%d'):
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                continue
        raise ValueError(f'Tanggal lahir "{value}" tidak dikenal (gunakan YYYY-MM-DD).')
