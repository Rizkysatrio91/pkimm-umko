"""Hapus massal BiodataKader hasil upload Excel (untuk uji LOKAL).

Membaca daftar NPM dari file Excel yang SAMA dengan yang Anda upload, lalu
menghapus BiodataKader yang NPM-nya cocok — sekali jalan, bukan satu per satu.

Default AMAN: hanya menghapus biodata TANPA akun user (id_user NULL), persis
seperti yang dibuat oleh fitur "Upload Kader via Excel". Dengan begitu akun /
biodata ASLI (yang punya user, mis. akun Anda sendiri) TIDAK ikut terhapus
walau NPM-nya kebetulan ada di file.

Pakai:
    python manage.py hapus_kader_dari_excel _uji_upload_kader/data_kader_uji_1.xlsx --dry-run
    python manage.py hapus_kader_dari_excel _uji_upload_kader/data_kader_uji_1.xlsx
    python manage.py hapus_kader_dari_excel _uji_upload_kader/data_kader_uji_2.xlsx --termasuk-berakun
"""
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from apps.kader.models import BiodataKader
from apps.pendaftaran.models import PendaftaranDAD


class Command(BaseCommand):
    help = (
        'Hapus massal BiodataKader berdasarkan NPM di file Excel upload. '
        'Default hanya yang tanpa akun user. Selalu jalankan --dry-run dulu.'
    )

    def add_arguments(self, parser):
        parser.add_argument('excel_path', type=str,
                            help='Path file Excel upload kader.')
        parser.add_argument('--dry-run', action='store_true',
                            help='Tampilkan apa yang AKAN dihapus tanpa menghapus.')
        parser.add_argument('--termasuk-berakun', action='store_true',
                            help='Ikut hapus biodata yang punya akun user '
                                 '(HATI-HATI: bisa data asli). Default: dilewati.')

    # ── Baca NPM dari Excel (kolom pertama, lewati header/contoh/petunjuk) ──
    def _baca_npm(self, path):
        p = Path(path)
        if not p.exists():
            raise CommandError(f'File tidak ditemukan: {path}')
        try:
            wb = load_workbook(p, data_only=True, read_only=True)
        except Exception as e:
            raise CommandError(f'Gagal membaca Excel {path}: {e}')

        ws = wb.active
        seen, npms = set(), []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            npm = str(row[0]).strip()
            if npm.endswith('.0'):          # angka Excel kadang jadi float
                npm = npm[:-2]
            if npm.isdigit() and len(npm) >= 7 and npm not in seen:
                seen.add(npm)
                npms.append(npm)
        wb.close()
        return npms

    def handle(self, *args, **opt):
        npms = self._baca_npm(opt['excel_path'])
        if not npms:
            self.stdout.write(self.style.WARNING(
                'Tidak ada NPM valid di file. Tidak ada yang dihapus.'))
            return

        qs = BiodataKader.objects.filter(npm__in=npms)
        tanpa_akun = qs.filter(id_user__isnull=True)
        berakun = qs.exclude(id_user__isnull=True)

        npm_dgn_pendaftaran = set(
            PendaftaranDAD.objects.filter(id_biodata__npm__in=npms)
            .values_list('id_biodata__npm', flat=True)
        )

        self.stdout.write(f'NPM valid di file  : {len(npms)}')
        self.stdout.write(f'Cocok di database  : {qs.count()}')
        self.stdout.write(f'  - tanpa akun     : {tanpa_akun.count()}  (target hapus)')
        self.stdout.write(
            f'  - punya akun     : {berakun.count()}'
            + ('  (IKUT dihapus)' if opt['termasuk_berakun']
               else '  (DILEWATI — pakai --termasuk-berakun bila perlu)'))
        if npm_dgn_pendaftaran:
            self.stdout.write(self.style.WARNING(
                f'  ! {len(npm_dgn_pendaftaran)} biodata punya PendaftaranDAD — '
                f'pendaftaran & nilai terkait ikut terhapus (cascade).'))
        if berakun.exists() and not opt['termasuk_berakun']:
            for b in berakun[:20]:
                self.stdout.write(f'    dilewati (punya akun): {b.npm} - {b.nama_lengkap}')

        target = qs if opt['termasuk_berakun'] else tanpa_akun
        n_target = target.count()

        if opt['dry_run']:
            self.stdout.write(self.style.NOTICE(
                f'[DRY-RUN] Akan menghapus {n_target} biodata. '
                f'Tidak ada yang dihapus.'))
            return

        if n_target == 0:
            self.stdout.write(self.style.WARNING(
                'Tidak ada biodata yang memenuhi syarat untuk dihapus.'))
            return

        with transaction.atomic():
            deleted = target.delete()
        self.stdout.write(self.style.SUCCESS(
            f'Selesai. Objek terhapus (termasuk relasi cascade): {deleted}'))
